from io import BytesIO

from docx import Document

from app.knowledge.parsers import parse_document


def test_parse_txt_and_markdown() -> None:
    assert parse_document("a.txt", b"hello") == "hello"
    assert parse_document("a.md", b"# Heading\nBody") == "# Heading\nBody"


def test_parse_docx() -> None:
    stream = BytesIO()
    document = Document()
    document.add_paragraph("Requirement one")
    document.add_paragraph("Requirement two")
    document.save(stream)
    text = parse_document("requirements.docx", stream.getvalue())
    assert "Requirement one" in text
    assert "Requirement two" in text
